from Bio import SeqIO
from Bio.Seq import Seq
from Bio import pairwise2
from Bio.pairwise2 import format_alignment
from Bio.Align import PairwiseAligner
import os
from docx import Document
from docx.shared import RGBColor, Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_LINE_SPACING
import re
from fpdf import FPDF
import numpy as np
from concurrent.futures import ThreadPoolExecutor
import multiprocessing
from Bio.SeqUtils import GC, molecular_weight
from Bio.SeqUtils.ProtParam import ProteinAnalysis

class AdvancedSequenceAnalyzer:
    def __init__(self):
        self.reference_seq = None
        self.reference_protein = None
        self.start_seq = None
        self.end_seq = None
        self.quality_threshold = 20
        self.cpu_count = multiprocessing.cpu_count()
        
    def read_multiple_formats(self, file_path):
        """Read multiple sequence file formats"""
        file_ext = os.path.splitext(file_path)[1].lower()
        
        try:
            if file_ext == '.ab1':
                return SeqIO.read(file_path, "abi")
            elif file_ext in ['.fasta', '.fa', '.fas']:
                return SeqIO.read(file_path, "fasta")
            elif file_ext in ['.fastq', '.fq']:
                return SeqIO.read(file_path, "fastq")
            elif file_ext in ['.gb', '.gbk']:
                return SeqIO.read(file_path, "genbank")
            elif file_ext == '.embl':
                return SeqIO.read(file_path, "embl")
            elif file_ext == '.txt':
                with open(file_path, 'r') as f:
                    content = f.read().strip()
                    # Try to detect if it's a sequence
                    clean_content = re.sub(r'[^ATGCN]', '', content.upper())
                    if len(clean_content) > len(content) * 0.8:  # 80% nucleotides
                        from Bio.SeqRecord import SeqRecord
                        return SeqRecord(Seq(clean_content), id="sequence")
                    else:
                        raise Exception("File doesn't appear to contain sequence data")
            else:
                raise Exception(f"Unsupported file format: {file_ext}")
        except Exception as e:
            raise Exception(f"Error reading file {file_path}: {str(e)}")

    def quality_filter_sequence(self, record, threshold=None):
        """Filter sequence based on quality scores"""
        if threshold is None:
            threshold = self.quality_threshold
            
        if not hasattr(record, 'letter_annotations') or 'phred_quality' not in record.letter_annotations:
            return str(record.seq)  # No quality data, return full sequence
            
        quality_scores = record.letter_annotations['phred_quality']
        sequence = str(record.seq)
        
        # Filter positions with quality below threshold
        filtered_seq = ""
        for i, (base, quality) in enumerate(zip(sequence, quality_scores)):
            if quality >= threshold:
                filtered_seq += base
            else:
                filtered_seq += 'N'  # Replace low quality bases with N
                
        return filtered_seq

    def advanced_trim_sequence(self, record, quality_threshold=20, window_size=10):
        """Advanced sequence trimming based on quality scores"""
        sequence = str(record.seq)
        
        if not hasattr(record, 'letter_annotations') or 'phred_quality' not in record.letter_annotations:
            return sequence  # No quality data, return full sequence
            
        quality_scores = record.letter_annotations['phred_quality']
        
        # Find start position
        start_pos = 0
        for i in range(len(quality_scores) - window_size):
            window_avg = np.mean(quality_scores[i:i + window_size])
            if window_avg >= quality_threshold:
                start_pos = i
                break
                
        # Find end position
        end_pos = len(quality_scores)
        for i in range(len(quality_scores) - window_size, window_size, -1):
            window_avg = np.mean(quality_scores[i - window_size:i])
            if window_avg >= quality_threshold:
                end_pos = i
                break
                
        return sequence[start_pos:end_pos]

    def parallel_alignment(self, sequences, reference_seq):
        """Perform alignments in parallel"""
        def align_single(seq_data):
            name, sequence = seq_data
            alignments = pairwise2.align.globalms(sequence, reference_seq, 2, -1, -10, -0.5)
            if alignments:
                return name, alignments[0]
            return name, None
            
        with ThreadPoolExecutor(max_workers=self.cpu_count) as executor:
            results = list(executor.map(align_single, sequences.items()))
            
        return dict(results)

    def calculate_sequence_statistics(self, sequence):
        """Calculate comprehensive sequence statistics"""
        sequence = sequence.upper()
        length = len(sequence)
        
        # Base composition
        base_counts = {
            'A': sequence.count('A'),
            'T': sequence.count('T'),
            'G': sequence.count('G'),
            'C': sequence.count('C'),
            'N': sequence.count('N')
        }
        
        base_percentages = {base: (count / length) * 100 for base, count in base_counts.items()}
        
        # GC content
        gc_content = GC(sequence)
        
        # Molecular weight
        mol_weight = molecular_weight(sequence, seq_type='DNA')
        
        # Dinucleotide frequencies
        dinucleotides = {}
        for i in range(len(sequence) - 1):
            dinuc = sequence[i:i+2]
            if 'N' not in dinuc:
                dinucleotides[dinuc] = dinucleotides.get(dinuc, 0) + 1
                
        return {
            'length': length,
            'base_counts': base_counts,
            'base_percentages': base_percentages,
            'gc_content': gc_content,
            'molecular_weight': mol_weight,
            'dinucleotide_frequencies': dinucleotides
        }

    def find_restriction_sites(self, sequence, enzymes=None):
        """Find restriction enzyme sites"""
        if enzymes is None:
            # Common restriction enzymes
            enzymes = {
                'EcoRI': 'GAATTC',
                'BamHI': 'GGATCC',
                'HindIII': 'AAGCTT',
                'XbaI': 'TCTAGA',
                'SacI': 'GAGCTC',
                'KpnI': 'GGTACC',
                'SmaI': 'CCCGGG',
                'PstI': 'CTGCAG'
            }
            
        sites = {}
        sequence = sequence.upper()
        
        for enzyme, site in enzymes.items():
            positions = []
            start = 0
            while True:
                pos = sequence.find(site, start)
                if pos == -1:
                    break
                positions.append(pos + 1)  # 1-based indexing
                start = pos + 1
            sites[enzyme] = positions
            
        return sites

    def translate_all_frames(self, sequence):
        """Translate sequence in all 6 reading frames"""
        sequence = sequence.upper()
        translations = {}
        
        # Forward frames
        for frame in range(3):
            frame_seq = sequence[frame:]
            if len(frame_seq) >= 3:
                protein = str(Seq(frame_seq).translate())
                translations[f'Frame +{frame + 1}'] = protein
                
        # Reverse frames
        rev_seq = str(Seq(sequence).reverse_complement())
        for frame in range(3):
            frame_seq = rev_seq[frame:]
            if len(frame_seq) >= 3:
                protein = str(Seq(frame_seq).translate())
                translations[f'Frame -{frame + 1}'] = protein
                
        return translations

    def advanced_alignment_analysis(self, seq1, seq2):
        """Perform advanced alignment analysis"""
        # Use PairwiseAligner for more control
        aligner = PairwiseAligner()
        aligner.match_score = 2
        aligner.mismatch_score = -1
        aligner.open_gap_score = -10
        aligner.extend_gap_score = -0.5
        
        alignments = aligner.align(seq1, seq2)
        best_alignment = alignments[0]
        
        # Calculate detailed statistics
        aligned_seq1 = str(best_alignment).split('\n')[0]
        aligned_seq2 = str(best_alignment).split('\n')[2]
        
        matches = sum(1 for a, b in zip(aligned_seq1, aligned_seq2) if a == b and a != '-')
        mismatches = sum(1 for a, b in zip(aligned_seq1, aligned_seq2) if a != b and a != '-' and b != '-')
        gaps = aligned_seq1.count('-') + aligned_seq2.count('-')
        
        identity = (matches / len(aligned_seq1)) * 100
        similarity = identity  # For DNA sequences
        
        return {
            'alignment': best_alignment,
            'score': best_alignment.score,
            'matches': matches,
            'mismatches': mismatches,
            'gaps': gaps,
            'identity': identity,
            'similarity': similarity,
            'aligned_length': len(aligned_seq1)
        }

    def export_advanced_results(self, results, output_path, format='pdf'):
        """Export results with advanced formatting"""
        if format.lower() == 'pdf':
            self._export_advanced_pdf(results, output_path)
        elif format.lower() == 'html':
            self._export_html(results, output_path)
        else:
            raise Exception(f"Unsupported export format: {format}")

    def _export_advanced_pdf(self, results, output_path):
        """Export to PDF with advanced formatting"""
        from fpdf import FPDF
        
        class AdvancedPDF(FPDF):
            def header(self):
                self.set_font('Arial', 'B', 15)
                self.cell(0, 10, 'SeqAnalyse - Advanced Analysis Results', 0, 1, 'C')
                self.ln(10)
                
            def footer(self):
                self.set_y(-15)
                self.set_font('Arial', 'I', 8)
                self.cell(0, 10, f'Page {self.page_no()}', 0, 0, 'C')
                
        pdf = AdvancedPDF()
        pdf.add_page()
        
        # Add results with better formatting
        for analysis_type, samples in results.items():
            pdf.set_font('Arial', 'B', 14)
            pdf.cell(0, 10, f'{analysis_type.capitalize()} Analysis', 0, 1)
            pdf.ln(5)
            
            for sample_name, result in samples.items():
                pdf.set_font('Arial', 'B', 12)
                pdf.cell(0, 8, f'Sample: {sample_name}', 0, 1)
                
                pdf.set_font('Courier', '', 8)
                lines = result.split('\n')
                for line in lines:
                    if len(line) > 80:
                        # Split long lines
                        for i in range(0, len(line), 80):
                            pdf.cell(0, 4, line[i:i+80], 0, 1)
                    else:
                        pdf.cell(0, 4, line, 0, 1)
                pdf.ln(5)
                
        pdf.output(output_path)

    def _export_html(self, results, output_path):
        """Export results to HTML format"""
        html_content = """
        <!DOCTYPE html>
        <html>
        <head>
            <title>SeqAnalyse Results</title>
            <style>
                body { font-family: Arial, sans-serif; margin: 20px; }
                .header { background-color: #f0f0f0; padding: 10px; text-align: center; }
                .analysis-section { margin: 20px 0; border: 1px solid #ccc; padding: 15px; }
                .sample { margin: 10px 0; }
                .sequence { font-family: 'Courier New', monospace; font-size: 12px; 
                           background-color: #f9f9f9; padding: 10px; white-space: pre-wrap; }
                .stats { background-color: #e9f4ff; padding: 10px; margin: 10px 0; }
            </style>
        </head>
        <body>
            <div class="header">
                <h1>SeqAnalyse - Advanced Analysis Results</h1>
            </div>
        """
        
        for analysis_type, samples in results.items():
            html_content += f'<div class="analysis-section"><h2>{analysis_type.capitalize()} Analysis</h2>'
            
            for sample_name, result in samples.items():
                html_content += f'<div class="sample"><h3>Sample: {sample_name}</h3>'
                html_content += f'<div class="sequence">{result}</div></div>'
                
            html_content += '</div>'
            
        html_content += '</body></html>'
        
        with open(output_path, 'w') as f:
            f.write(html_content)